#!/usr/bin/env python3
"""把地址映射报告里 unmatched 的旧地址导成 Excel，供人工逐条确认。

只做导出和归因，不猜房：每条给出旧库原文、目标小区、以及没匹配上的原因分类，
再附上 PMS 侧该小区现有楼栋/房号清单，便于人工回填 house_id 后再导入。
"""

from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# 明显不是住人房号的关键词，这类旧地址本来就不该落到某一间房上
NON_HOUSE_WORDS = ("停车", "车位", "办公", "青客", "仓库", "门面", "商铺", "公共", "宿舍")


def read_csv(path: Path):
    with open(path, "rt", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def read_pms_rooms(path: Path):
    """community_id -> {building_no -> set(room_no)}；顺带记住小区名。"""
    buildings: dict[str, dict[str, set[str]]] = defaultdict(lambda: defaultdict(set))
    names: dict[str, str] = {}
    with open(path, "rt", encoding="utf-8") as stream:
        for line in stream:
            line = line.rstrip("\n")
            if not line.strip():
                continue
            community_id, community_name, building_no, room_no = line.split("|")
            buildings[community_id][building_no].add(room_no)
            names[community_id] = community_name
    return buildings, names


def guess_buildings(raw_building: str, pms: dict[str, set[str]]):
    """旧库写「A栋」、PMS 写「A栋34号」，按前缀找同名楼栋。只给候选，不自动采用。"""
    if not raw_building:
        return []
    exact = [name for name in pms if name == raw_building]
    if exact:
        return exact
    return sorted(name for name in pms if name.startswith(raw_building))


def classify(row: dict, buildings: dict[str, dict[str, set[str]]]):
    """返回（原因分类, 说明）。只解释，不替人决定对应哪间房。"""
    community_id = (row.get("target_community_id") or "").strip()
    raw_room = (row.get("房间号") or "").strip()
    raw_building = (row.get("楼栋号") or "").strip()
    pms = buildings.get(community_id) or {}

    if not pms:
        return "PMS 该小区尚未建房产", "PMS 里这个小区还没有任何楼栋/房间，需先建档再匹配"
    if any(word in raw_room or word in raw_building for word in NON_HOUSE_WORDS):
        return "非住人房号", "旧地址写的是停车位/办公室这类非房号，不应绑定到具体房间"
    if not raw_room:
        return "旧地址缺房号", "旧库房间号为空，无法定位"

    candidates = guess_buildings(raw_building, pms)
    plain_room = re.sub(r"[（(].*?[)）]", "", raw_room).strip()
    has_code = plain_room != raw_room

    if not candidates:
        return "楼栋在 PMS 不存在", f"PMS 该小区楼栋为：{'、'.join(sorted(pms))}"
    if len(candidates) > 1:
        return "楼栋写法对应多个候选", f"旧「{raw_building}」可能是：{'、'.join(candidates)}，需人工指定"

    building = candidates[0]
    hint = f"旧「{raw_building}」对应 PMS「{building}」" if building != raw_building else f"PMS「{building}」"
    if plain_room in pms[building]:
        if has_code:
            return "房号含旧定位码（去码后能对上）", f"{hint}，去掉括号后房号 {plain_room} 在 PMS 存在，待人工确认是否同一间"
        return "待复核：房号其实能对上", f"{hint} 下存在房号 {plain_room}，原报告判为未匹配，需复核匹配规则"
    if has_code:
        return "房号含旧定位码（去码后仍对不上）", f"{hint} 下没有房号 {plain_room}"
    return "房号在该楼栋不存在", f"{hint} 下没有这个房号"


def style_header(sheet, widths):
    fill = PatternFill("solid", fgColor="31558A")
    for index, width in enumerate(widths, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--mapping", type=Path, required=True)
    parser.add_argument("--pms-rooms", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()

    rows = [r for r in read_csv(args.mapping) if r.get("match_status") == "unmatched"]
    buildings, names = read_pms_rooms(args.pms_rooms)

    book = Workbook()
    sheet = book.active
    sheet.title = "未匹配地址"
    headers = [
        "旧地址ID", "旧小区编号", "旧楼栋号", "旧房间号", "旧库姓名", "旧库手机号",
        "目标PMS小区ID", "目标PMS小区", "没匹配上的原因", "说明",
        "人工确认房号(待填)", "人工确认PMS房屋ID(待填)", "备注(待填)",
    ]
    sheet.append(headers)
    for row in rows:
        reason, note = classify(row, buildings)
        sheet.append([
            row.get("ID", ""), row.get("小区编号", ""), row.get("楼栋号", ""), row.get("房间号", ""),
            row.get("姓名", ""), row.get("手机号", ""),
            row.get("target_community_id", ""), row.get("target_community", ""),
            reason, note, "", "", "",
        ])
    style_header(sheet, [10, 11, 12, 20, 12, 14, 14, 18, 20, 46, 18, 20, 20])

    summary = book.create_sheet("原因汇总")
    summary.append(["没匹配上的原因", "目标PMS小区", "条数"])
    counter: dict[tuple[str, str], int] = defaultdict(int)
    for row in rows:
        reason, _ = classify(row, buildings)
        counter[(reason, row.get("target_community", "") or "未映射")] += 1
    for (reason, community), count in sorted(counter.items(), key=lambda item: -item[1]):
        summary.append([reason, community, count])
    summary.append(["合计", "", len(rows)])
    style_header(summary, [24, 20, 8])

    reference = book.create_sheet("PMS现有楼栋房号")
    reference.append(["PMS小区ID", "PMS小区", "楼栋号", "房间数", "房号清单"])
    for community_id in sorted(buildings, key=int):
        for building_no, room_nos in sorted(buildings[community_id].items()):
            ordered = sorted(room_nos, key=lambda value: (len(value), value))
            reference.append([
                community_id, names.get(community_id, ""), building_no,
                len(ordered), "、".join(ordered),
            ])
    style_header(reference, [12, 20, 14, 10, 120])

    args.out.parent.mkdir(parents=True, exist_ok=True)
    book.save(args.out)
    print(f"{args.out} 共 {len(rows)} 条未匹配地址")


if __name__ == "__main__":
    main()
